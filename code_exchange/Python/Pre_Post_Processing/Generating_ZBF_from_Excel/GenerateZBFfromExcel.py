import pandas as pd
from openpyxl import load_workbook
from tkinter import filedialog
import os

# Author: Yuan Chen
# Date: 2025/01/26
# It's written to convert Excel data into ZBF text format. A sample Excel is provided as ZBF_Example.xlsx.
# The generated ZBF file has the same name and is placed in the same folder as the Excel file.
# After generation, please copy the ZBF file to your \Documents\Zemax\POP\BEAMFILES folder. Alternatively you can change line 16 to your current BEAMFILES folder.

# Read Excel name
file_path = filedialog.askopenfilename(title="Select Excel File", filetypes=[("Excel files", "*.xlsx")])
file_folder = os.path.dirname(file_path)
file_name = os.path.splitext(os.path.basename(file_path))[0]
output_zbf_path = os.path.join(file_folder, f"{file_name}.zbf")

# Extract basic data of ZBF, the sampling and the polarization
tabs = ['ZBF', 'Ex_RealPart', 'Ex_ImaginaryPart']
workbook = load_workbook(file_path, data_only=True)
data_frames = []
first_tab = workbook[tabs[0]]
a1_to_a30_data = [first_tab[f'A{i}'].value for i in range(1, 31)]
X_sampling = first_tab['A3'].value
Y_sampling = first_tab['A4'].value
Ispol = first_tab['A5'].value

if Ispol:
    tabs = ['ZBF', 'Ex_RealPart', 'Ex_ImaginaryPart', 'Ey_RealPart', 'Ey_ImaginaryPart']

# Extract the ZBF information
data_frames.append(pd.DataFrame(a1_to_a30_data, columns=['A1_to_A30']))
for tab_name in tabs[1:]:
    sheet = workbook[tab_name]
    tab_data = [[cell.value for cell in row] for row in sheet.iter_rows()]
    data_frames.append(pd.DataFrame(tab_data))

# Check Data
for idx, tab in enumerate(tabs[1:], start=1):
    print(f'Original data size for {tab}: {data_frames[idx].shape}')

# Only extract data with the sampling size
for idx in range(1, len(data_frames)):
    data_frames[idx] = data_frames[idx].iloc[:X_sampling, :Y_sampling]

# Output the size of each array
for idx, tab in enumerate(tabs[1:], start=1):
    rows, cols = data_frames[idx].shape
    print(f'After slicing, data size for {tab}: {rows} rows, {cols} columns')

# Create the ZBF
with open(output_zbf_path, 'w') as file:
    # Information needed by ZBF before the electric field data
    for value in data_frames[0]['A1_to_A30']:
        file.write(f"{value}\n")
    # Ex data
    for i in range(X_sampling):
        for j in range(Y_sampling):
            # Ex_real and Ex_imaginary
            data_A = data_frames[1].iloc[i, j]
            data_B = data_frames[2].iloc[i, j]
            file.write(f'{data_A}\n')
            file.write(f'{data_B}\n')
    if Ispol:
        for i in range(X_sampling):
            for j in range(Y_sampling):
                # Ey_real and Ey_imaginary
                data_C = data_frames[3].iloc[i, j]
                data_D = data_frames[4].iloc[i, j]
                file.write(f'{data_C}\n')
                file.write(f'{data_D}\n')
print("Data have been written as ZBF.")
